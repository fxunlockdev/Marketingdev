"""Synthetic scanned guideline and composed slides for real extraction tests."""
import io
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.utils import ImageReader
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

out = Path("work/document-fixtures")
out.mkdir(parents=True, exist_ok=True)
image = Image.new("RGB", (1200, 1600), "#faf8ef")
draw = ImageDraw.Draw(image)
font_path = next((path for path in ["/System/Library/Fonts/Supplemental/Arial.ttf",
    "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"] if Path(path).is_file()), None)
if font_path is None:
    raise RuntimeError("Arial or Liberation Sans fixture font required")
font = ImageFont.truetype(font_path, 38)
for index, line in enumerate(["Brand name: Olive and Field", "Offering: Botanical soap", "Voice: Clear and warm", "Audience: Homeware shops"]):
    draw.text((80, 140 + index * 100), line, fill="#233c31", font=font)
draw.rectangle((80, 650, 1120, 1150), fill="#2355c4")
draw.text((120, 800), "REFILL. REUSE. REPEAT.", fill="white", font=font)
canvas = Canvas(str(out / "scanned-guidelines.pdf"), pagesize=(600, 800))
canvas.drawImage(ImageReader(image), 0, 0, width=600, height=800)
canvas.showPage()
canvas.save()
presentation = Presentation()
for index in range(2):
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(35, 85, 196)
    title = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(8), Inches(1))
    title.text_frame.text = "Olive and Field" if index == 0 else "Hidden brand guideline"
    run = title.text_frame.paragraphs[0].runs[0]
    run.font.size = Pt(38)
    run.font.color.rgb = RGBColor(255, 255, 255)
    stream = io.BytesIO()
    image.save(stream, "PNG")
    stream.seek(0)
    slide.shapes.add_picture(stream, Inches(0.7), Inches(1.8), height=Inches(4.7))
    if index == 1:
        slide._element.set("show", "0")
presentation.save(out / "composed-guidelines.pptx")
print("Created synthetic scanned PDF and two composed slides, including a hidden slide")
